"""
Comprehensive API Integration Test for Resume Upload & Conversion Endpoints.
"""

import os
import sys

# Ensure backend directory is in sys.path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "backend")))

from fastapi.testclient import TestClient
from app.main import app

client = TestClient(app)


def test_root_status():
    """Verify system status includes /api/resume."""
    res = client.get("/")
    assert res.status_code == 200
    data = res.json()
    assert "/api/resume" in data["endpoints"]


def test_upload_rejects_disguised_exe():
    """Verify that an exe renamed to .pdf is rejected with 400 and the required error message."""
    fake_exe = b"MZ" + b"\x00" * 200
    files = {"file": ("malicious.pdf", fake_exe, "application/pdf")}
    res = client.post("/api/resume/upload", files=files)
    assert res.status_code == 400
    body = res.json()
    assert "content does not match" in body.get("message", "")


def test_upload_rejects_oversized_file():
    """Verify that files > 10MB are rejected with 400 and the required error message."""
    huge_payload = b"%PDF-1.7\n" + b"0" * (10 * 1024 * 1024 + 100)
    files = {"file": ("big_resume.pdf", huge_payload, "application/pdf")}
    res = client.post("/api/resume/upload", files=files)
    assert res.status_code == 400
    body = res.json()
    assert "File is too large" in body.get("message", "")


def _build_sample_resume_docx() -> bytes:
    """Build a minimal but realistic resume .docx in-memory so this test
    doesn't depend on a pre-generated file in the gitignored storage/ dir."""
    import io
    from docx import Document

    doc = Document()
    doc.add_heading("Gaurav Pratap", level=1)
    doc.add_paragraph("Email: gaurav.pratap@example.com | Phone: +91-9876543210")
    doc.add_heading("Summary", level=2)
    doc.add_paragraph("Cloud engineer with 5 years of experience in Azure DevOps and CI/CD pipelines.")
    doc.add_heading("Experience", level=2)
    doc.add_paragraph("Senior DevOps Engineer, TechCorp (2021-Present)")
    doc.add_paragraph("Led migration of build pipelines to Azure DevOps, reducing deployment time by 40%.")
    doc.add_heading("Skills", level=2)
    doc.add_paragraph("Azure DevOps, Kubernetes, Terraform, Python, CI/CD")
    doc.add_heading("Education", level=2)
    doc.add_paragraph("B.Tech in Computer Science, XYZ University (2017)")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_upload_and_convert_flow_with_docx():
    """Verify full end-to-end upload, DLP detection, AI structuring, LaTeX, and DOCX generation."""
    file_bytes = _build_sample_resume_docx()

    # 1. Upload
    files = {"file": ("Gaurav_Resume.docx", file_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    res = client.post("/api/resume/upload", files=files)
    assert res.status_code == 200, f"Upload failed: {res.text}"
    upload_res = res.json()
    assert upload_res["status"] == "success"
    upload_id = upload_res["upload_id"]
    assert upload_id is not None
    assert upload_res["detected_type"] == "DOCX"

    # 2. Convert
    convert_payload = {"upload_id": upload_id}
    res_conv = client.post("/api/resume/convert", json=convert_payload)
    assert res_conv.status_code == 200, f"Convert failed: {res_conv.text}"
    conv_data = res_conv.json()
    assert conv_data["status"] == "ready"
    task_id = conv_data["task_id"]
    assert task_id is not None
    assert conv_data["completeness_report"]["preservation_score"] >= 95.0

    # 3. Download DOCX
    res_docx = client.get(f"/api/resume/download/{task_id}/docx")
    assert res_docx.status_code == 200
    assert len(res_docx.content) > 5000
    assert "openxmlformats" in res_docx.headers.get("content-type", "")

    # 4. Download LaTeX .tex
    res_tex = client.get(f"/api/resume/download/{task_id}/tex")
    assert res_tex.status_code == 200
    assert r"\documentclass{resume}" in res_tex.text

    # 5. Preview endpoint
    res_prev = client.get(f"/api/resume/preview/{task_id}")
    assert res_prev.status_code == 200
    prev_data = res_prev.json()
    assert "structured_resume" in prev_data
    assert "latex_code" in prev_data


if __name__ == "__main__":
    test_root_status()
    print("test_root_status passed!")
    test_upload_rejects_disguised_exe()
    print("test_upload_rejects_disguised_exe passed!")
    test_upload_rejects_oversized_file()
    print("test_upload_rejects_oversized_file passed!")
    test_upload_and_convert_flow_with_docx()
    print("test_upload_and_convert_flow_with_docx passed!")
    print("\nALL API INTEGRATION TESTS PASSED SUCCESSFULLY!")
