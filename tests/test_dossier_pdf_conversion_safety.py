"""
Regression tests for the Dossier Compiler's pdf2docx corruption detection.

Real bug: for PDFs like the fixture below (LaTeX-generated), the third-party
pdf2docx library — used by DossierService.compile_dossier_docx() to embed an
"exact" copy of a PDF resume — silently drops inter-word spaces ("Senior
Security Analyst" -> "SeniorSecurityAnalyst") and can drop whole sections
(a small skills table vanished entirely). Total character count barely
changes in the missing-spaces case (only space characters are lost), so the
pre-existing length-based safety net (_resume_embedded_ok's character-count
check) does not catch it. _docx_avg_word_length is a second, independent
signal for the same corruption check: real prose averages ~4-7 characters
per word; glued-together text averages 20-40+.
"""

import io
import os
import sys

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "backend")))

import docx
from app.services.dossier_service import DossierService

FIXTURE_PATH = os.path.join(os.path.dirname(__file__), "fixtures", "single_column_resume_with_skills_table.pdf")


def _empty_docx_snapshot() -> bytes:
    buf = io.BytesIO()
    docx.Document().save(buf)
    return buf.getvalue()


def test_pdf2docx_actually_corrupts_this_real_pdf_confirming_the_bug_exists():
    """Sanity-checks the bug this fix targets is real, not hypothetical —
    run pdf2docx directly (bypassing the app) against the real fixture PDF."""
    from pdf2docx import Converter

    out_path = "/tmp/_test_pdf2docx_corruption_check.docx"
    cv = Converter(FIXTURE_PATH)
    cv.convert(out_path)
    cv.close()

    d = docx.Document(out_path)
    words = []
    for p in d.paragraphs:
        words.extend(p.text.split())

    avg_len = sum(len(w) for w in words) / len(words)
    assert avg_len > 18, (
        "This test fixture is expected to trigger the pdf2docx space-loss bug "
        "(that's WHY it's used as a fixture) — if this now fails, pdf2docx's "
        "behavior may have changed and _resume_embedded_ok's threshold should "
        "be re-validated, not just have this assertion loosened."
    )


def test_resume_embedded_ok_rejects_space_corrupted_conversion_output():
    svc = DossierService()
    with open(FIXTURE_PATH, "rb") as f:
        pdf_bytes = f.read()

    corrupted_bytes = svc.convert_pdf_to_docx_bytes(pdf_bytes)
    assert corrupted_bytes, "convert_pdf_to_docx_bytes must succeed (pdf2docx runs fine, it just corrupts the text)"

    corrupted_path = "/tmp/_test_corrupted_dossier_output.docx"
    with open(corrupted_path, "wb") as f:
        f.write(corrupted_bytes)

    assert svc._resume_embedded_ok(corrupted_path, pdf_bytes, _empty_docx_snapshot()) is False


def test_docx_avg_word_length_distinguishes_corrupted_from_normal_text():
    svc = DossierService()

    normal_doc = docx.Document()
    normal_doc.add_paragraph(
        "Senior Security Analyst with fourteen years of experience in data protection, "
        "data loss prevention, and access governance across multiple large organizations."
    )
    normal_path = "/tmp/_test_normal_word_length.docx"
    normal_doc.save(normal_path)
    assert svc._docx_avg_word_length(normal_path) < 10

    corrupted_words = [
        "SeniorSecurityAnalystwithfourteenyearsofexperienceindataprotection",
        "AndAnotherReallyLongGluedTogetherRunOnWordLikeThisOneRightHere",
        "OneMoreForGoodMeasureToBeSafe",
        "SomeMoreWordsHereForStatisticalConfidence",
        "YetAnotherAbsurdlyLongConcatenatedTokenToPadOutTheSampleSize",
        "AndOneFinalGluedTogetherPhraseToClearTheFifteenWordMinimum",
        "PlusAFewExtraJustInCaseTheThresholdNeedsMoreMargin",
        "AbsolutelyPositivelyOneMoreLongRunOnWordHereToo",
    ] * 2  # 16 tokens total, comfortably over the 15-token confidence floor
    corrupted_doc = docx.Document()
    corrupted_doc.add_paragraph(" ".join(corrupted_words))
    corrupted_path = "/tmp/_test_corrupted_word_length.docx"
    corrupted_doc.save(corrupted_path)
    assert svc._docx_avg_word_length(corrupted_path) > 18


def test_docx_avg_word_length_returns_none_for_too_little_text_to_judge():
    svc = DossierService()
    sparse_doc = docx.Document()
    sparse_doc.add_paragraph("Short text.")
    sparse_path = "/tmp/_test_sparse_word_length.docx"
    sparse_doc.save(sparse_path)
    assert svc._docx_avg_word_length(sparse_path) is None
