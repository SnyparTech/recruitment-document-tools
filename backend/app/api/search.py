import io
import logging
import os
import re
import time
from typing import List, Optional
from fastapi import APIRouter, File, HTTPException, UploadFile, status
from pydantic import BaseModel, Field
import fitz  # PyMuPDF
import docx

from app.schemas.requirement import (
    CandidateSearchRequest,
    CandidateSearchResponse,
    ExecutionResult,
    ValidationResult,
)
from app.schemas.search_plan import SearchPlan
from app.schemas.candidate_result import CandidateResult, SubmitCandidateResultsRequest
from app.services.requirement_service import RequirementService
from app.services.candidate_store_service import merge_candidates
from app.services.candidate_ranking_service import rank_candidates
from app.services import state_persistence

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/search", tags=["Candidate Search"])

requirement_service = RequirementService()

ALLOWED_DOC_EXTENSIONS = {".pdf", ".docx", ".doc", ".txt"}
MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB


# NOTE: same single-tenant, process-global storage pattern documented below —
# known limitation (see audit P0/P1 follow-up on session isolation), not
# addressed in this change. What IS addressed: these globals are now backed
# by a JSON file (state_persistence) so a backend restart no longer silently
# loses the recruiter's in-progress plan/candidates.
_persisted = state_persistence.load_state()
_latest_search_plan = _persisted.get("search_plan")
_latest_plan_timestamp = _persisted.get("plan_timestamp") or 0.0
_latest_candidates: List[CandidateResult] = [
    CandidateResult(**c) for c in (_persisted.get("candidates") or [])
]
_latest_candidates_timestamp = _persisted.get("candidates_timestamp") or 0.0
_latest_search_id: Optional[str] = _persisted.get("search_id")

# Requirement-chat draft state: a SEPARATE plan from _latest_search_plan
# above. The chat edits this one turn-by-turn; nothing here reaches the
# extension until the recruiter clicks "Apply to Resdex" on the frontend,
# which POSTs this draft to /search/plan (store_search_plan) like any other
# pre-built plan. Keeps "still drafting" cleanly separate from "live on Resdex".
_draft_plan: Optional[dict] = _persisted.get("draft_plan")
_draft_timestamp: float = _persisted.get("draft_timestamp") or 0.0
_chat_history: List[dict] = _persisted.get("chat_history") or []


def _persist_now() -> None:
    state_persistence.save_state(
        search_plan=_latest_search_plan,
        plan_timestamp=_latest_plan_timestamp,
        candidates=[c.model_dump() for c in _latest_candidates],
        candidates_timestamp=_latest_candidates_timestamp,
        search_id=_latest_search_id,
        draft_plan=_draft_plan,
        draft_timestamp=_draft_timestamp,
        chat_history=_chat_history,
    )


@router.get(
    "/active-plan",
    summary="Get active SearchPlan for Naukri Resdex browser auto-fill",
)
async def get_active_plan():
    """Returns the latest SearchPlan for the extension to auto-fill."""
    global _latest_search_plan, _latest_plan_timestamp
    return {
        "status": "success",
        "has_plan": _latest_search_plan is not None,
        "timestamp": _latest_plan_timestamp,
        "plan": _latest_search_plan,
    }


@router.post(
    "/candidates",
    response_model=CandidateSearchResponse,
    status_code=status.HTTP_200_OK,
    summary="Generate SearchPlan from natural language requirement",
)
async def search_candidates(
    request: CandidateSearchRequest,
) -> CandidateSearchResponse:
    """
    Parses natural language requirement into structured SearchPlan.
    Plan is stored for the extension to pick up.
    """
    global _latest_search_plan, _latest_plan_timestamp, _latest_candidates, _latest_candidates_timestamp

    plan, validation = requirement_service.process_requirement(request.requirement)

    if request.active_in:
        plan.active_in = request.active_in
    if request.verified_mobile is not None:
        plan.verified_mobile = request.verified_mobile
    if request.verified_email is not None:
        plan.verified_email = request.verified_email
    if request.attached_resume is not None:
        plan.attached_resume = request.attached_resume

    # Re-validate after the request-level overrides above, since active_in
    # can be supplied directly by the caller and bypass the agent's own checks.
    validation = requirement_service.validator.validate(plan)
    if not validation.valid:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail={
                "message": "Generated SearchPlan failed Resdex schema validation and was not stored.",
                "errors": validation.errors,
                "warnings": validation.warnings,
            },
        )

    _latest_candidates = []
    _latest_candidates_timestamp = time.time()
    plan_dict = plan.model_dump()
    plan_dict["_submit_search"] = request.submit_search
    _latest_search_plan = plan_dict
    _latest_plan_timestamp = time.time()
    _persist_now()

    msg = "SearchPlan generated. Extension will auto-fill the form."
    if request.submit_search:
        msg = "SearchPlan generated. Extension will auto-fill AND submit search."

    execution_result = ExecutionResult(
        requested=request.submit_search,
        executed=False,
        form_filled=False,
        search_submitted=False,
        fields_interacted=[],
        message=msg,
    )

    return CandidateSearchResponse(
        requirement=request.requirement,
        search_plan=plan,
        validation=validation,
        execution=execution_result,
    )


class DirectSearchPlanRequest(BaseModel):
    plan: SearchPlan = Field(..., description="Pre-built SearchPlan to store for extension")
    submit_search: bool = Field(default=False, description="Whether extension should auto-click Search Candidates")


@router.post(
    "/plan",
    response_model=CandidateSearchResponse,
    status_code=status.HTTP_200_OK,
    summary="Store a pre-built SearchPlan for extension",
)
async def store_search_plan(
    request: DirectSearchPlanRequest,
) -> CandidateSearchResponse:
    """Stores a pre-built SearchPlan for the extension to pick up, after Resdex schema validation."""
    global _latest_search_plan, _latest_plan_timestamp

    plan = request.plan

    validation = requirement_service.validator.validate(plan)
    if not validation.valid:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail={
                "message": "Submitted SearchPlan failed Resdex schema validation and was not stored.",
                "errors": validation.errors,
                "warnings": validation.warnings,
            },
        )

    plan_dict = plan.model_dump()
    plan_dict["_submit_search"] = request.submit_search
    _latest_search_plan = plan_dict
    _latest_plan_timestamp = time.time()
    _persist_now()

    msg = "SearchPlan stored. Extension will auto-fill the form."
    if request.submit_search:
        msg = "SearchPlan stored. Extension will auto-fill AND submit search."

    execution_result = ExecutionResult(
        requested=request.submit_search,
        executed=False,
        form_filled=False,
        search_submitted=False,
        fields_interacted=[],
        message=msg,
    )

    return CandidateSearchResponse(
        requirement="(direct SearchPlan)",
        search_plan=plan,
        validation=validation,
        execution=execution_result,
    )


@router.post(
    "/results",
    status_code=status.HTTP_200_OK,
    summary="Submit candidates scraped from the Resdex results page",
)
async def submit_candidate_results(request: SubmitCandidateResultsRequest):
    """
    Receives a batch of candidates scraped by the extension from the current
    Resdex results page, deduplicates against previously stored candidates,
    and stores the merged list. Extraction diagnostics (container/field hit
    counts, no candidate content) are logged for selector calibration but not
    persisted.
    """
    global _latest_candidates, _latest_candidates_timestamp, _latest_search_id

    # A new Resdex search (e.g. user pressed Modify) replaces the old list.
    if request.search_id and request.search_id != _latest_search_id:
        _latest_candidates = []
        _latest_search_id = request.search_id

    if request.diagnostics:
        logger.info(
            "Candidate extraction diagnostics (page %s): containers=%s extracted=%s "
            "strategy=%s field_hits=%s warnings=%s",
            request.page,
            request.diagnostics.containers_detected,
            request.diagnostics.candidates_extracted,
            request.diagnostics.selector_strategy,
            request.diagnostics.field_hit_counts,
            request.diagnostics.warnings,
        )

    if not request.candidates:
        _persist_now()  # search_id/reset above may have changed state even with an empty batch
        return {
            "status": "success",
            "message": "No candidates in this batch (nothing to merge).",
            "total_stored": len(_latest_candidates),
        }

    _latest_candidates = merge_candidates(_latest_candidates, request.candidates)
    _latest_candidates_timestamp = time.time()
    _persist_now()

    return {
        "status": "success",
        "message": f"Merged {len(request.candidates)} candidate(s) from page {request.page}.",
        "total_stored": len(_latest_candidates),
    }


@router.get(
    "/results",
    summary="Get all candidates extracted so far for the active search, ranked against the active SearchPlan",
)
async def get_candidate_results():
    """
    Returns all deduplicated candidates scraped so far. If a SearchPlan is
    currently active, each candidate is ranked against it (explainable score +
    data completeness, see candidate_ranking_service.py); otherwise candidates
    are returned unscored rather than ranked against nothing.
    """
    active_plan = SearchPlan(**_latest_search_plan) if _latest_search_plan else None
    ranked = rank_candidates(_latest_candidates, active_plan)
    return {
        "status": "success",
        "has_results": len(_latest_candidates) > 0,
        "timestamp": _latest_candidates_timestamp,
        "count": len(_latest_candidates),
        "ranked_against_active_plan": active_plan is not None,
        "candidates": [c.model_dump() for c in ranked],
    }


@router.delete(
    "/results",
    summary="Clear stored candidate results (e.g. when starting a new search)",
)
async def clear_candidate_results():
    global _latest_candidates, _latest_candidates_timestamp
    _latest_candidates = []
    _latest_candidates_timestamp = 0.0
    _persist_now()
    return {"status": "success", "message": "Candidate results cleared."}


class UpdateKeywordMandatoryRequest(BaseModel):
    required: List[str] = Field(..., description="Keywords HR wants marked mandatory (starred) in Resdex")
    preferred: List[str] = Field(default_factory=list, description="Remaining tracked keywords, not mandatory")


@router.patch(
    "/plan/keywords",
    status_code=status.HTTP_200_OK,
    summary="Update which keywords are mandatory on the active SearchPlan, and re-apply live on Resdex",
)
async def update_plan_keyword_mandatory(request: UpdateKeywordMandatoryRequest):
    """
    Lets HR interactively pick a SUBSET of the plan's keywords as mandatory
    (rather than all-or-nothing) after a plan already exists. Moves keywords
    between keywords.required/preferred on the active plan, bumps the plan
    timestamp, and flags `_keyword_sync_only` so the extension re-applies
    just the star toggles on the live Resdex tab (via "Modify" + a fresh
    search) instead of re-typing the whole form.
    """
    global _latest_search_plan, _latest_plan_timestamp

    if not _latest_search_plan:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="No active SearchPlan to update. Generate a search first.",
        )

    kw = _latest_search_plan.get("keywords") or {}
    kw["required"] = request.required
    kw["preferred"] = request.preferred
    _latest_search_plan["keywords"] = kw
    _latest_search_plan["_submit_search"] = True
    _latest_search_plan["_keyword_sync_only"] = True
    _latest_plan_timestamp = time.time()
    _persist_now()

    return {
        "status": "success",
        "message": f"{len(request.required)} keyword(s) marked mandatory. Extension will re-apply on Resdex.",
        "plan": _latest_search_plan,
    }


class ChatEditRequest(BaseModel):
    message: str = Field(..., min_length=1, description="First message = a JD/requirement; later messages = edit instructions")


@router.get(
    "/plan/chat",
    summary="Get the requirement chat history and current (unapplied) draft SearchPlan",
)
async def get_chat_state():
    """Restores the chat thread + draft plan on page refresh. The draft is
    separate from the applied plan (/active-plan) — this never reflects what's
    live on Resdex, only what the recruiter has staged so far."""
    return {
        "status": "success",
        "history": _chat_history,
        "plan": _draft_plan,
        "timestamp": _draft_timestamp,
    }


@router.post(
    "/plan/chat",
    summary="Chat-edit the draft SearchPlan (paste a JD to start, or send follow-up edit instructions)",
)
async def chat_edit_plan(request: ChatEditRequest):
    """
    Requirement chat, staged: nothing here touches the live Resdex tab. The
    first message (no draft yet) is parsed as a full JD; every message after
    is applied as an edit on top of the existing draft. The recruiter reviews
    the resulting plan + keyword pills on the frontend and explicitly clicks
    "Apply to Resdex" (POST /search/plan) to actually push it to the extension.
    """
    global _draft_plan, _draft_timestamp, _chat_history

    base_plan = SearchPlan(**_draft_plan) if _draft_plan else None
    updated_plan, reply = requirement_service.chat_edit(base_plan, request.message, _chat_history)

    _chat_history = _chat_history + [
        {"role": "user", "content": request.message},
        {"role": "assistant", "content": reply},
    ]
    _draft_plan = updated_plan.model_dump()
    _draft_timestamp = time.time()
    _persist_now()

    return {
        "status": "success",
        "reply": reply,
        "plan": _draft_plan,
        "history": _chat_history,
    }


@router.delete(
    "/plan/chat",
    summary="Clear the requirement chat and draft SearchPlan (start over)",
)
async def clear_chat_state():
    global _draft_plan, _draft_timestamp, _chat_history
    _draft_plan = None
    _draft_timestamp = 0.0
    _chat_history = []
    _persist_now()
    return {"status": "success", "message": "Chat and draft plan cleared."}


@router.post(
    "/upload-requirement",
    status_code=status.HTTP_200_OK,
    summary="Extract text from Requirement Document / PDF",
)
async def upload_requirement_document(
    file: UploadFile = File(..., description="Requirement or Job Description document (PDF, DOCX, DOC, TXT)"),
):
    """Extracts raw, structured requirement text from uploaded document."""
    ext = os.path.splitext(file.filename or "")[1].lower()
    if ext not in ALLOWED_DOC_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Unsupported file format '{ext}'. Allowed formats: {', '.join(sorted(ALLOWED_DOC_EXTENSIONS))}",
        )

    file_bytes = await file.read()
    if len(file_bytes) == 0:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Uploaded file is empty.")
    if len(file_bytes) > MAX_FILE_SIZE_BYTES:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="File exceeds maximum allowed size of 10 MB.")

    extracted_text = ""
    try:
        if ext == ".pdf":
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            page_texts = []
            for page in doc:
                text = page.get_text()
                if text.strip():
                    page_texts.append(text)
            extracted_text = "\n\n".join(page_texts).strip()

        elif ext == ".docx":
            doc_stream = io.BytesIO(file_bytes)
            docx_doc = docx.Document(doc_stream)
            paragraphs = [p.text for p in docx_doc.paragraphs if p.text.strip()]
            table_lines = []
            for table in docx_doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        table_lines.append(" | ".join(cells))
            extracted_text = "\n".join(paragraphs)
            if table_lines:
                extracted_text += "\n\n" + "\n".join(table_lines)

        elif ext == ".txt":
            try:
                extracted_text = file_bytes.decode("utf-8")
            except UnicodeDecodeError:
                extracted_text = file_bytes.decode("latin-1", errors="replace")

        elif ext == ".doc":
            text_chars = re.findall(rb"[\x20-\x7E\r\n\t]{4,}", file_bytes)
            extracted_text = "\n".join(tc.decode("latin-1", errors="ignore") for tc in text_chars)

    except Exception as exc:
        logger.error(f"Failed to extract requirement text from {file.filename}: {exc}")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to read requirement document: {str(exc)}",
        )

    if not extracted_text.strip():
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Could not extract readable text from the document. Please ensure it is not an empty or scanned image file.",
        )

    return {
        "status": "success",
        "filename": file.filename,
        "size": len(file_bytes),
        "extracted_text": extracted_text.strip(),
        "char_count": len(extracted_text.strip()),
        "line_count": len(extracted_text.strip().splitlines()),
    }
